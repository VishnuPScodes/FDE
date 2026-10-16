from io import BytesIO
from openpyxl import load_workbook

def read_workbook(data):
    workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    iterator = sheet.iter_rows(values_only=True)
    headers = [str(value) for value in next(iterator)]
    return [dict(zip(headers, row)) for row in iterator if any(value is not None for value in row)]
